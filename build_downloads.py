import os
import zipfile
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import markdown_it

BASE_DIR = r"C:\Users\harie\.gemini\antigravity\scratch\eles"
REPORT_MD_PATH = os.path.join(BASE_DIR, "ELES_PROJECT_REPORT.md")
DOCX_OUTPUT_PATH = os.path.join(BASE_DIR, "ELES_Project_Report.docx")
HTML_OUTPUT_PATH = os.path.join(BASE_DIR, "ELES_Project_Report.html")
ZIP_OUTPUT_PATH = os.path.join(BASE_DIR, "ELES_Complete_Project.zip")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_docx():
    print("Generating DOCX report...")
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    with open(REPORT_MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    in_table = False
    table_rows = []

    for line in lines:
        raw = line.rstrip('\n')
        stripped = raw.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(raw)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            # Check if divider row
            if re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', stripped):
                continue
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            # End of table, write table to docx
            if table_rows:
                cols_count = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=cols_count)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(table_rows):
                    for c_idx, val in enumerate(row):
                        if c_idx < cols_count:
                            cell = t.cell(r_idx, c_idx)
                            cell.text = val
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            for run in p.runs:
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(10)
                                if r_idx == 0:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            if r_idx == 0:
                                set_cell_background(cell, '1E293B')
                            elif r_idx % 2 == 1:
                                set_cell_background(cell, 'F8FAFC')
            table_rows = []
            in_table = False

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', stripped[2:])
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        else:
            # Clean HTML tags if any
            clean_text = re.sub(r'<[^>]+>', '', stripped)
            if clean_text:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
                parts = re.split(r'(\*\*.*?\*\*)', clean_text)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

    doc.save(DOCX_OUTPUT_PATH)
    print(f"DOCX report saved to {DOCX_OUTPUT_PATH}")

def generate_html():
    print("Generating HTML report...")
    with open(REPORT_MD_PATH, 'r', encoding='utf-8') as f:
        md_content = f.read()

    md = markdown_it.MarkdownIt("commonmark", {"breaks": True, "html": True})
    html_body = md.render(md_content)

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>ELES Project Report - Kongu Engineering College</title>
    <link href="https://fonts.googleapis.com/css2?family=Times+New+Roman&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
    <style>
        body {{
            font-family: 'Times New Roman', Times, serif;
            font-size: 14pt;
            line-height: 1.6;
            color: #111827;
            background-color: #f3f4f6;
            margin: 0;
            padding: 20px;
        }}
        .report-container {{
            max-width: 900px;
            margin: 0 auto;
            background: #ffffff;
            padding: 60px 80px;
            box-shadow: 0 4px 20px rgba(0,0,0,0.08);
            border-radius: 8px;
        }}
        .top-action-bar {{
            position: sticky;
            top: 20px;
            display: flex;
            justify-content: flex-end;
            gap: 12px;
            margin-bottom: 24px;
            z-index: 100;
        }}
        .btn {{
            background: #4f46e5;
            color: #ffffff;
            border: none;
            padding: 10px 20px;
            font-size: 14px;
            font-weight: bold;
            font-family: system-ui, sans-serif;
            border-radius: 6px;
            cursor: pointer;
            box-shadow: 0 2px 8px rgba(79, 70, 229, 0.3);
            text-decoration: none;
            display: inline-flex;
            align-items: center;
            gap: 8px;
        }}
        .btn-secondary {{
            background: #059669;
            box-shadow: 0 2px 8px rgba(5, 150, 105, 0.3);
        }}
        .btn:hover {{
            filter: brightness(1.1);
        }}
        h1, h2, h3, h4 {{
            font-family: 'Times New Roman', Times, serif;
            color: #0f172a;
            page-break-after: avoid;
        }}
        h1 {{
            text-align: center;
            font-size: 20pt;
            margin-top: 30px;
            margin-bottom: 15px;
            text-transform: uppercase;
        }}
        h2 {{
            font-size: 16pt;
            border-bottom: 2px solid #e2e8f0;
            padding-bottom: 6px;
            margin-top: 30px;
        }}
        h3 {{
            font-size: 14pt;
            margin-top: 20px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            font-size: 11pt;
        }}
        th, td {{
            border: 1px solid #cbd5e1;
            padding: 8px 12px;
            text-align: left;
        }}
        th {{
            background-color: #1e293b;
            color: #ffffff;
            font-weight: bold;
        }}
        tr:nth-child(even) {{
            background-color: #f8fafc;
        }}
        code, pre {{
            font-family: 'JetBrains Mono', monospace;
            background: #f1f5f9;
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 10pt;
        }}
        pre {{
            padding: 16px;
            overflow-x: auto;
            border: 1px solid #e2e8f0;
            line-height: 1.4;
        }}
        @media print {{
            body {{
                background: #ffffff;
                padding: 0;
            }}
            .report-container {{
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .top-action-bar {{
                display: none;
            }}
            table, pre {{
                page-break-inside: avoid;
            }}
            h1, h2, h3 {{
                page-break-after: avoid;
            }}
        }}
    </style>
</head>
<body>
    <div class="top-action-bar">
        <button class="btn" onclick="window.print()">🖨️ Print / Save as PDF</button>
        <a class="btn btn-secondary" href="/download/report.docx">📥 Download Word (.docx)</a>
        <a class="btn" href="/download/project.zip">📦 Download Project (.zip)</a>
    </div>

    <div class="report-container">
        {html_body}
    </div>
</body>
</html>
"""
    with open(HTML_OUTPUT_PATH, 'w', encoding='utf-8') as f:
        f.write(full_html)
    print(f"HTML report saved to {HTML_OUTPUT_PATH}")

def generate_zip():
    print("Generating project ZIP...")
    exclude_dirs = {'node_modules', '.pytest_cache', '__pycache__', '.git', 'dist'}
    with zipfile.ZipFile(ZIP_OUTPUT_PATH, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, dirs, files in os.walk(BASE_DIR):
            dirs[:] = [d for d in dirs if d not in exclude_dirs]
            for file in files:
                if file.endswith('.zip'):
                    continue
                file_path = os.path.join(root, file)
                rel_path = os.path.relpath(file_path, BASE_DIR)
                zipf.write(file_path, arcname=rel_path)
    print(f"ZIP package saved to {ZIP_OUTPUT_PATH}")

if __name__ == "__main__":
    generate_docx()
    generate_html()
    generate_zip()
    print("All downloads built successfully!")
